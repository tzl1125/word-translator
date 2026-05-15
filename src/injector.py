"""
注入器模块：将翻译结果写回到新的Word文档中。
"""

import json
import os
import shutil
import tempfile
import zipfile
from typing import List, Dict
from lxml import etree
from docx import Document
from docx.oxml.shared import qn
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_UNDERLINE, WD_BREAK
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph

from .core import NAMESPACES


class Injector:
    """
    翻译结果注入器。

    读取包含翻译结果的检查点文件，并将所有内容（包括文本、格式、
    图表、SmartArt等）精确地写入到一个新的.docx文件中。
    """

    def __init__(self, input_file: str, checkpoint_file: str, output_file: str):
        """
        初始化注入器。

        Args:
            input_file: 原始输入文件（用于复制基础结构）。
            checkpoint_file: 包含翻译结果的检查点文件。
            output_file: 最终输出的翻译后文件。
        """
        self.input_file = input_file
        self.checkpoint_file = checkpoint_file
        self.output_file = output_file
        self.doc = Document(input_file)

    def _apply_run(self, para: Paragraph, run_info: Dict):
        """将单个RunInfo应用到段落中。"""
        if run_info.get('is_page_break'):
            run = para.add_run()
            run.add_break(WD_BREAK.PAGE)
            return
        orig_xml = run_info.get('original_xml', '')
        is_complex = bool(orig_xml.strip()) and (
                run_info.get('is_field_code') or run_info.get('is_math') or run_info.get(
            'is_image') or run_info.get('is_toc_field'))
        if is_complex:
            try:
                elem = etree.fromstring(orig_xml)
                para._element.append(elem)
                return
            except Exception as e:
                print(f"警告：无法解析 original_xml，回退为文本。错误: {e}")
        text = run_info.get('translated_text', run_info.get('text', ''))
        if text:
            run = para.add_run(text)
            # 应用样式
            for attr in ['bold', 'italic', 'underline', 'superscript', 'subscript', 'strike']:
                if run_info.get(attr) is not None:
                    setattr(run.font, attr, run_info[attr])
            # 字体名称、大小、下划线类型
            rPr = run._element.get_or_add_rPr()
            if run_info.get('font_name') or run_info.get('font_name_east_asia') or run_info.get('font_name_cs'):
                rFonts = rPr.find(qn('w:rFonts')) or etree.SubElement(rPr, qn('w:rFonts'))
                for attr, tag in [('font_name', 'w:ascii'), ('font_name_east_asia', 'w:eastAsia'),
                                  ('font_name_cs', 'w:cs')]:
                    if run_info.get(attr):
                        rFonts.set(qn(tag), run_info[attr])
            if run_info.get('font_size'): run.font.size = Pt(run_info['font_size'])
            if run_info.get('underline_type'): run.font.underline = WD_UNDERLINE.SINGLE if run_info[
                                                                                               'underline_type'] == 1 else WD_UNDERLINE.NONE
            # 底纹和颜色
            if run_info.get('shading'):
                shd = rPr.find(qn('w:shd')) or etree.SubElement(rPr, qn('w:shd'))
                for k in ['val', 'color', 'fill']:
                    if run_info['shading'].get(k):
                        shd.set(qn(f'w:{k}'), run_info['shading'][k])
            if run_info.get('color'): run.font.color.rgb = RGBColor(*run_info['color'])
            if run_info.get('highlight'): run.font.highlight_color = run_info['highlight']

    def _apply_runs(self, para: Paragraph, runs_list: List[Dict]):
        """清空段落并应用新的run列表。"""
        # 清空段落除 pPr 外的所有子元素
        for child in list(para._element):
            if child.tag != qn('w:pPr'):
                para._element.remove(child)
        for r in runs_list:
            self._apply_run(para, r)

    def _clear_except_images(self, para: Paragraph):
        """清空段落，但保留图片。"""
        for child in list(para._element):
            if child.tag != qn('w:pPr') and child.find('.//w:drawing', NAMESPACES) is None:
                para._element.remove(child)

    def _inject_headers_footers(self, segments: List[Dict]):
        """注入页眉页脚的翻译结果。"""
        headers = {}
        footers = {}
        for seg in segments:
            d = headers if seg['is_header'] else footers
            d.setdefault(seg['section_idx'], []).append(seg)
        for si, sec in enumerate(self.doc.sections):
            for is_header, elem in [(True, sec.header), (False, sec.footer)]:
                if not elem: continue
                segs = headers.get(si, []) if is_header else footers.get(si, [])
                if not segs: continue
                para_elems = [e for e in elem._element.iter(qn('w:p')) if isinstance(e, CT_P)]
                for seg in sorted(segs, key=lambda x: x['para_idx']):
                    if seg['para_idx'] < len(para_elems):
                        p = Paragraph(para_elems[seg['para_idx']], elem)
                        self._apply_runs(p, seg['runs_list'])

    def _process_all_zip_modifications(self, chart_segments: List[Dict], smartart_segments: List[Dict]):
        """直接修改ZIP包内的XML文件以更新图表和SmartArt。"""
        if not chart_segments and not smartart_segments:
            return

        temp_dir = tempfile.mkdtemp()
        try:
            with zipfile.ZipFile(self.output_file, 'r') as z:
                z.extractall(temp_dir)

            chart_by_file = {}
            for seg in chart_segments:
                chart_by_file.setdefault(seg['file_path'], []).append(seg)

            for file_path, seg_list in chart_by_file.items():
                full_path = os.path.join(temp_dir, file_path)
                if not os.path.exists(full_path):
                    continue
                tree = etree.parse(full_path)
                root = tree.getroot()
                for seg in seg_list:
                    # 根据 element_type 定位节点（当前仅支持 'title'）
                    if seg['element_type'] == 'title':
                        titles = root.findall('.//c:title', NAMESPACES)
                        if seg['element_idx'] < len(titles):
                            texts = titles[seg['element_idx']].findall('.//a:t', NAMESPACES)
                            if texts:
                                texts[0].text = seg['translated_text']
                tree.write(full_path, encoding='UTF-8', xml_declaration=True)

            smartart_by_file = {}
            for seg in smartart_segments:
                smartart_by_file.setdefault(seg['file_path'], []).append(seg)

            for file_path, seg_list in smartart_by_file.items():
                full_path = os.path.join(temp_dir, file_path)
                if not os.path.exists(full_path):
                    continue
                tree = etree.parse(full_path)
                root = tree.getroot()
                for seg in seg_list:
                    elems = root.findall('.//a:t', NAMESPACES)
                    if seg['element_idx'] < len(elems):
                        elems[seg['element_idx']].text = seg['translated_text']
                tree.write(full_path, encoding='UTF-8', xml_declaration=True)

            # 重新打包为 DOCX
            with zipfile.ZipFile(self.output_file, 'w', zipfile.ZIP_DEFLATED) as z:
                for root_dir, _, files in os.walk(temp_dir):
                    for f in files:
                        full = os.path.join(root_dir, f)
                        z.write(full, os.path.relpath(full, temp_dir))
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def inject(self):
        """执行完整的注入流程。"""
        with open(self.checkpoint_file, encoding="utf-8") as f:
            data = json.load(f)

        # 重建段落列表（包括 sdt 内的目录）
        body = self.doc.element.body
        all_paras = []
        for child in body:
            if child.tag == qn('w:p'):
                all_paras.append(Paragraph(child, self.doc))
            elif child.tag == qn('w:sdt'):
                sdt_content = child.find(qn('w:sdtContent'))
                if sdt_content is not None:
                    for p_elem in sdt_content.findall(qn('w:p')):
                        all_paras.append(Paragraph(p_elem, self.doc))

        for seg in data["text_segments"]:
            idx = seg["seg_idx"]
            if idx >= len(all_paras):
                print(f"警告：段落索引 {idx} 超出范围")
                continue
            para = all_paras[idx]
            if seg.get("is_toc_paragraph", False):
                runs = seg["runs_list"]
                if len(runs) == 1 and runs[0].get("original_xml"):
                    try:
                        new_elem = etree.fromstring(runs[0]["original_xml"])
                        parent = para._element.getparent()
                        if parent is not None:
                            parent.replace(para._element, new_elem)
                            all_paras[idx] = Paragraph(new_elem, self.doc)
                        else:
                            para.clear()
                            para._element.append(new_elem)
                    except Exception as e:
                        print(f"替换目录段落失败: {e}")
                        para.clear()
                        self._apply_runs(para, runs)
                else:
                    para.clear()
                    self._apply_runs(para, runs)
                continue
            if seg["has_smartart_or_chart"]:
                self._clear_except_images(para)
            else:
                para.clear()
            self._apply_runs(para, seg["runs_list"])

        # 表格
        for seg in data["table_cell_segments"]:
            cell = self.doc.tables[seg['table_idx']].rows[seg['row_idx']].cells[seg['cell_idx']]
            para = cell.paragraphs[seg['para_idx']]
            para.clear()
            self._apply_runs(para, seg['runs_list'])

        # 页眉页脚
        self._inject_headers_footers(data["header_footer_segments"])

        # 设置更新域
        update = self.doc.settings._element.find(qn('w:updateFields'))
        if update is None:
            update = etree.SubElement(self.doc.settings._element, qn('w:updateFields'))
        update.set(qn('w:val'), 'true')

        self.doc.save(self.output_file)

        # 注入图表和 SmartArt
        self._process_all_zip_modifications(data["chart_segments"], data["smartart_segments"])
